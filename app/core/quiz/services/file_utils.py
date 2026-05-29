import hashlib
import logging
import os
import re
import tempfile
from io import BytesIO

import docx
import pytesseract
from PIL import Image, ImageFilter, ImageOps
from app.core.config import settings
from app.core.quiz.database.quizzes import delete_document, get_document_by_hash
from app.core.quiz.services.vectorize import extract_chapter_contents_from_bytes, extract_chapter_headings_with_page_numbers_from_bytes

logger = logging.getLogger(__name__)

try:
    import fitz
except ModuleNotFoundError:  # PyMuPDF is optional until PDF quiz parsing is used.
    fitz = None


def _require_fitz():
    if fitz is None:
        raise RuntimeError("PDF parsing requires PyMuPDF. Install the 'pymupdf' package to enable this feature.")
    return fitz

if os.getenv("TESSERACT_CMD"):
    pytesseract.pytesseract.tesseract_cmd = os.getenv("TESSERACT_CMD")

IMAGE_FILE_TYPES = {"jpg", "jpeg", "png", "webp", "tif", "tiff", "bmp"}
PARSER_VERSION = "ocr-v3"
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", str(25 * 1024 * 1024)))

def hash_content(content):
    if isinstance(content, str):
        content = content.encode("utf-8")
    return hashlib.sha256(content).hexdigest()

def hash_uploaded_file(content: bytes) -> str:
    return hash_content(PARSER_VERSION.encode("utf-8") + content)

def extract_author_from_docx(binary_data: bytes) -> str:
    try:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(binary_data)
            tmp_path = tmp.name
        doc = docx.Document(tmp_path)
        author = doc.core_properties.author
        return author if author else "Unknown Author"
    except Exception:
        return "Unknown Author"
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass

def extract_author_from_pdf(binary_data: bytes) -> str:
    try:
        fitz_lib = _require_fitz()
        pdf = fitz_lib.open(stream=binary_data, filetype="pdf")
        metadata = pdf.metadata
        return metadata.get("author") or "Unknown Author"
    except Exception:
        return "Unknown Author"

def extract_author_from_image(_binary_data: bytes) -> str:
    return "Unknown Author"

def extract_title_from_filename(filename: str) -> str:
    name = os.path.splitext(os.path.basename(filename))[0]
    return name.replace("_", " ").replace("-", " ").strip() or "Unknown Title"

def save_file_locally(doc_hash, file_ext, content):
    upload_dir = settings.data_dir / "quiz_uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{doc_hash}.{file_ext}"
    file_path = upload_dir / filename
    with open(file_path, "wb") as f:
        f.write(content)

def serialize_existing_chapters(chapters):
    return [
        {
            "id": ch["id"],
            "title": ch["chapter_title"],
            "chapter_title": ch["chapter_title"],
            "content": ch["content"],
            "hash": ch["chapter_hash"],
        }
        for ch in chapters
    ]

def process_file_content(filename: str, content: bytes):
    if len(content) > MAX_UPLOAD_BYTES:
        limit_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        raise ValueError(f"File is too large. Upload a file up to {limit_mb} MB.")

    file_ext = filename.split(".")[-1].lower()

    if file_ext == "pdf":
        doc_author = extract_author_from_pdf(content)
    elif file_ext == "docx":
        doc_author = extract_author_from_docx(content)
    elif file_ext in IMAGE_FILE_TYPES:
        doc_author = extract_author_from_image(content)
    else:
        raise ValueError("Unsupported file type")

    doc_title = extract_title_from_filename(filename)
    doc_hash = hash_uploaded_file(content)

    existing_doc = get_document_by_hash(doc_hash)
    if existing_doc:
        existing_chapters = serialize_existing_chapters(existing_doc.get("chapters", []))
        if existing_chapters:
            return {
                "document_id": existing_doc["id"],
                "already_exists": True,
                "document_hash": existing_doc["document_hash"],
                "title": existing_doc["title"],
                "author": existing_doc["author"],
                "file_type": existing_doc["file_type"],
                "chapters": existing_chapters,
            }
        # If no chapters, we might want to delete and reprocess, but let's just reprocess
        # delete_document(existing_doc["id"], existing_doc["owner_id"])

    chapters = []
    if file_ext == "pdf":
        chapter_headings = extract_chapter_headings_with_page_numbers_from_bytes(content)
        chapters = extract_chapter_contents_from_bytes(content, chapter_headings)

    if not chapters:
        text = extract_text_from_file(content, file_ext)
        chapters = split_into_chapters(text)

    if not chapters or (len(chapters) == 1 and len(chapters[0].get("content", "")) > 6000):
        text = extract_text_from_file(content, file_ext)
        chapters = split_into_study_sections(text)

    if not chapters:
        raise ValueError("No readable text found.")

    for index, ch in enumerate(chapters, start=1):
        ch["hash"] = hash_content(f"{doc_hash}:{index}:{ch['content']}")

    save_file_locally(doc_hash, file_ext, content)

    return {
        "already_exists": False,
        "document_hash": doc_hash,
        "title": doc_title,
        "author": doc_author,
        "file_type": file_ext,
        "chapters": chapters,
    }

def extract_text_from_pdf(binary_data: bytes) -> str:
    fitz_lib = _require_fitz()
    pdf = fitz_lib.open(stream=binary_data, filetype="pdf")
    text = clean_ocr_text("\n".join(page.get_text() for page in pdf))
    if len(text.strip()) >= 40:
        return text
    return extract_text_from_scanned_pdf(binary_data)

def extract_text_from_docx(binary_data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(binary_data)
        tmp_path = tmp.name
    try:
        doc = docx.Document(tmp_path)
        return clean_ocr_text("\n".join(p.text for p in doc.paragraphs))
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass

def extract_text_from_image(binary_data: bytes) -> str:
    try:
        image = Image.open(BytesIO(binary_data))
        return extract_best_ocr_text(image)
    except Exception as exc:
        raise RuntimeError("OCR failed") from exc

def extract_text_from_scanned_pdf(binary_data: bytes) -> str:
    fitz_lib = _require_fitz()
    pdf = fitz_lib.open(stream=binary_data, filetype="pdf")
    extracted_pages = []
    for page in pdf:
        pixmap = page.get_pixmap(matrix=fitz_lib.Matrix(3, 3), alpha=False)
        image = Image.open(BytesIO(pixmap.tobytes("png")))
        extracted_pages.append(extract_best_ocr_text(image))
    return clean_ocr_text("\n".join(extracted_pages))

def extract_text_from_file(binary_data: bytes, file_ext: str) -> str:
    if file_ext == "pdf":
        return extract_text_from_pdf(binary_data)
    if file_ext == "docx":
        return extract_text_from_docx(binary_data)
    if file_ext in IMAGE_FILE_TYPES:
        return extract_text_from_image(binary_data)
    raise ValueError("Unsupported file type")

def preprocess_ocr_image(image: Image.Image, threshold: int | None = None) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    image = image.convert("L")
    width, height = image.size
    scale = max(1, min(4, 2200 // max(width, 1)))
    if scale > 1:
        image = image.resize((width * scale, height * scale), Image.Resampling.LANCZOS)
    image = ImageOps.autocontrast(image)
    image = image.filter(ImageFilter.MedianFilter(size=3))
    if threshold is not None:
        image = image.point(lambda pixel: 255 if pixel > threshold else 0)
    return image

def ocr_quality_score(text: str) -> float:
    cleaned = clean_ocr_text(text)
    if len(cleaned) < 30:
        return 0
    chars = [char for char in cleaned if not char.isspace()]
    if not chars:
        return 0
    alpha_ratio = sum(char.isalpha() for char in chars) / len(chars)
    symbol_ratio = sum((not char.isalnum()) and char not in ".,:;?!()-/%" for char in chars) / len(chars)
    words = re.findall(r"[A-Za-z][A-Za-z]{2,}", cleaned)
    return (alpha_ratio * 2) + min(len(words) / 80, 1) - (symbol_ratio * 1.5)

def clean_ocr_text(text: str) -> str:
    text = text.replace("\x0c", "\n")
    text = re.sub(r"[|_~`]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    lines = []
    for raw_line in re.split(r"(?<=[.!?])\s+|\n+", text):
        line = raw_line.strip(" -•*")
        if not line:
            continue
        compact = line.replace(" ", "")
        if len(compact) < 8:
            continue
        alpha_ratio = sum(char.isalpha() for char in compact) / max(len(compact), 1)
        if alpha_ratio < 0.45:
            continue
        if len(re.findall(r"[~`{}_^\\|]", line)) > 1:
            continue
        lines.append(line)
    return " ".join(lines).strip()

def extract_best_ocr_text(image: Image.Image) -> str:
    variants = [
        preprocess_ocr_image(image),
        preprocess_ocr_image(image, threshold=155),
        preprocess_ocr_image(image, threshold=180),
    ]
    configs = [
        "--oem 3 --psm 6",
        "--oem 3 --psm 4",
        "--oem 3 --psm 11",
    ]
    candidates = []
    for variant in variants:
        for config in configs:
            try:
                text = pytesseract.image_to_string(variant, config=config)
                candidates.append((ocr_quality_score(text), clean_ocr_text(text)))
            except:
                candidates.append((0, ""))
    best_score, best_text = max(candidates, key=lambda candidate: candidate[0])
    if best_score < 0.7:
        raise ValueError("OCR text quality too low")
    return best_text

def split_into_chapters(text: str):
    pattern = re.compile(r"\b(Chapter|Lesson|Unit)\s+([0-9IVXLC]+)\b", re.IGNORECASE)
    matches = list(pattern.finditer(text))
    chapters = []
    for i, match in enumerate(matches):
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        section_type = match.group(1).capitalize()
        section_number = match.group(2)
        title = f"{section_type} {section_number}"
        content = text[start:end].strip()
        chapters.append({"title": title, "content": content, "section_type": section_type})
    return chapters

def split_into_study_sections(text: str, words_per_section: int = 900):
    words = text.split()
    if not words:
        return []
    sections = []
    for index, start in enumerate(range(0, len(words), words_per_section), start=1):
        section_words = words[start:start + words_per_section]
        content = " ".join(section_words).strip()
        if content:
            sections.append({
                "title": f"Study Section {index}",
                "content": content,
                "section_type": "Auto",
            })
    return sections

def build_chapters_from_text(text: str, document_hash: str = "edited"):
    cleaned = clean_ocr_text(text)
    chapters = split_into_chapters(cleaned)
    if not chapters or (len(chapters) == 1 and len(chapters[0].get("content", "")) > 6000):
        chapters = split_into_study_sections(cleaned, words_per_section=650)
    if not chapters:
        raise ValueError("No usable study text found.")
    for index, chapter in enumerate(chapters, start=1):
        chapter["hash"] = hash_content(f"{document_hash}:edited:{index}:{chapter['content']}")
    return chapters
