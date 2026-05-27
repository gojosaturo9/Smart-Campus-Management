import hashlib
import logging
import os
import re
import tempfile
from io import BytesIO

import docx
import fitz
import pytesseract
from PIL import Image, ImageFilter, ImageOps
from dotenv import load_dotenv
from app.crud import delete_document, get_document_by_hash
from sqlalchemy.orm import Session
from app.vectorize import extract_chapter_contents_from_bytes, extract_chapter_headings_with_page_numbers_from_bytes

logger = logging.getLogger(__name__)

load_dotenv()
if os.getenv("TESSERACT_CMD"):
    pytesseract.pytesseract.tesseract_cmd = os.getenv("TESSERACT_CMD")

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

IMAGE_FILE_TYPES = {"jpg", "jpeg", "png", "webp", "tif", "tiff", "bmp"}
PARSER_VERSION = "ocr-v3"
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", str(25 * 1024 * 1024)))

def hash_content(content):
    if isinstance(content, str):
        content = content.encode("utf-8")
    return hashlib.sha256(content).hexdigest()

def hash_uploaded_file(content: bytes) -> str:
    return hash_content(PARSER_VERSION.encode("utf-8") + content)

def extract_author_from_docx(binary_data: bytes) -> str:
    """Extract author from DOCX core properties, falling back to 'Unknown Author'."""
    try:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(binary_data)
            tmp_path = tmp.name
        doc = docx.Document(tmp_path)
        author = doc.core_properties.author
        return author if author else "Unknown Author"
    except Exception:
        return "Unknown Author"
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass

def extract_author_from_pdf(binary_data: bytes) -> str:
    """Extract author from PDF metadata, falling back to 'Unknown Author'."""
    try:
        pdf = fitz.open(stream=binary_data, filetype="pdf")
        metadata = pdf.metadata
        return metadata.get("author") or "Unknown Author"
    except Exception:
        return "Unknown Author"

def extract_author_from_image(_binary_data: bytes) -> str:
    return "Unknown Author"

def extract_title_from_filename(filename: str) -> str:
    """Extract a human-readable title from the uploaded filename."""
    name = os.path.splitext(os.path.basename(filename))[0]
    return name.replace("_", " ").replace("-", " ").strip() or "Unknown Title"

def save_file(doc_hash, file_ext, content):
    filename = f"{doc_hash}.{file_ext}"
    file_path = os.path.join(UPLOAD_DIR, filename)
    with open(file_path, "wb") as f:
        f.write(content)

def serialize_existing_chapters(chapters):
    return [
        {
            "id": chapter.id,
            "title": chapter.chapter_title,
            "chapter_title": chapter.chapter_title,
            "content": chapter.content,
            "hash": chapter.chapter_hash,
        }
        for chapter in chapters
    ]

async def process_file(file, db: Session):
    content = await file.read()
    return process_file_content(file.filename, content, db)

def process_file_content(filename: str, content: bytes, db: Session):
    if len(content) > MAX_UPLOAD_BYTES:
        limit_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        raise ValueError(f"File is too large. Upload a file up to {limit_mb} MB.")

    file_ext = filename.split(".")[-1].lower()

    if file_ext == "pdf":
        doc_author = extract_author_from_pdf(content)
    elif file_ext == "docx":
        doc_author = extract_author_from_docx(content)
    elif file_ext in IMAGE_FILE_TYPES:
        doc_author = extract_author_from_image(content)
    else:
        raise ValueError("Unsupported file type")

    doc_title = extract_title_from_filename(filename)
    doc_hash = hash_uploaded_file(content)

    # Use content hash for reliable duplicate detection
    existing_doc = get_document_by_hash(db, doc_hash)
    if existing_doc:
        existing_chapters = serialize_existing_chapters(existing_doc.chapters)
        if existing_chapters:
            logger.info("Document already exists: id=%s", existing_doc.id)
            return {
                "document_id": existing_doc.id,
                "already_exists": True,
                "document_hash": existing_doc.document_hash,
                "title": existing_doc.title,
                "author": existing_doc.author,
                "file_type": existing_doc.file_type,
                "chapters": existing_chapters,
            }
        logger.warning("Cached document id=%s has no chapters; reprocessing upload.", existing_doc.id)
        delete_document(db, existing_doc)

    chapters = []
    if file_ext == "pdf":
        chapter_headings = extract_chapter_headings_with_page_numbers_from_bytes(content)
        chapters = extract_chapter_contents_from_bytes(content, chapter_headings)

    if not chapters:
        text = extract_text_from_file(content, file_ext)
        chapters = split_into_chapters(text)

    if not chapters or (len(chapters) == 1 and len(chapters[0].get("content", "")) > 6000):
        text = extract_text_from_file(content, file_ext)
        chapters = split_into_study_sections(text)

    if not chapters:
        raise ValueError(
            "No readable text was found. For scanned files or photos, install Tesseract OCR "
            "and upload a clear image/PDF."
        )

    logger.info("Extracted %d chapters from %s", len(chapters), filename)

    for index, ch in enumerate(chapters, start=1):
        ch["hash"] = hash_content(f"{doc_hash}:{index}:{ch['content']}")

    save_file(doc_hash, file_ext, content)

    return {
        "already_exists": False,
        "document_hash": doc_hash,
        "title": doc_title,
        "author": doc_author,
        "file_type": file_ext,
        "chapters": chapters,
    }

def extract_text_from_pdf(binary_data: bytes) -> str:
    """Extract plain text from a PDF file."""
    pdf = fitz.open(stream=binary_data, filetype="pdf")
    text = clean_ocr_text("\n".join(page.get_text() for page in pdf))
    if len(text.strip()) >= 40:
        return text
    return extract_text_from_scanned_pdf(binary_data)

def extract_text_from_docx(binary_data: bytes) -> str:
    """Extract plain text from a DOCX file."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(binary_data)
        tmp_path = tmp.name
    try:
        doc = docx.Document(tmp_path)
        return clean_ocr_text("\n".join(p.text for p in doc.paragraphs))
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass

def extract_text_from_image(binary_data: bytes) -> str:
    """Extract text from an image using local Tesseract OCR."""
    try:
        image = Image.open(BytesIO(binary_data))
        return extract_best_ocr_text(image)
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(
            "Tesseract OCR is not installed or not available in PATH. "
            "Install Tesseract to read photos and scanned files."
        ) from exc

def extract_text_from_scanned_pdf(binary_data: bytes) -> str:
    """OCR PDF pages when normal PDF text extraction returns little/no text."""
    try:
        pdf = fitz.open(stream=binary_data, filetype="pdf")
        extracted_pages = []
        for page in pdf:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
            image = Image.open(BytesIO(pixmap.tobytes("png")))
            extracted_pages.append(extract_best_ocr_text(image))
        return clean_ocr_text("\n".join(extracted_pages))
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(
            "Tesseract OCR is not installed or not available in PATH. "
            "Install Tesseract to read scanned PDFs."
        ) from exc

def extract_text_from_file(binary_data: bytes, file_ext: str) -> str:
    if file_ext == "pdf":
        return extract_text_from_pdf(binary_data)
    if file_ext == "docx":
        return extract_text_from_docx(binary_data)
    if file_ext in IMAGE_FILE_TYPES:
        return extract_text_from_image(binary_data)
    raise ValueError("Unsupported file type")

def preprocess_ocr_image(image: Image.Image, threshold: int | None = None) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    image = image.convert("L")
    width, height = image.size
    scale = max(1, min(4, 2200 // max(width, 1)))
    if scale > 1:
        image = image.resize((width * scale, height * scale), Image.Resampling.LANCZOS)
    image = ImageOps.autocontrast(image)
    image = image.filter(ImageFilter.MedianFilter(size=3))
    if threshold is not None:
        image = image.point(lambda pixel: 255 if pixel > threshold else 0)
    return image

def ocr_quality_score(text: str) -> float:
    cleaned = clean_ocr_text(text)
    if len(cleaned) < 30:
        return 0
    chars = [char for char in cleaned if not char.isspace()]
    if not chars:
        return 0
    alpha_ratio = sum(char.isalpha() for char in chars) / len(chars)
    symbol_ratio = sum((not char.isalnum()) and char not in ".,:;?!()-/%" for char in chars) / len(chars)
    words = re.findall(r"[A-Za-z][A-Za-z]{2,}", cleaned)
    return (alpha_ratio * 2) + min(len(words) / 80, 1) - (symbol_ratio * 1.5)

def clean_ocr_text(text: str) -> str:
    text = text.replace("\x0c", "\n")
    text = re.sub(r"[|_~`]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    lines = []
    for raw_line in re.split(r"(?<=[.!?])\s+|\n+", text):
        line = raw_line.strip(" -•*")
        if not line:
            continue
        compact = line.replace(" ", "")
        if len(compact) < 8:
            continue
        alpha_ratio = sum(char.isalpha() for char in compact) / max(len(compact), 1)
        if alpha_ratio < 0.45:
            continue
        if len(re.findall(r"[~`{}_^\\|]", line)) > 1:
            continue
        lines.append(line)
    return " ".join(lines).strip()

def extract_best_ocr_text(image: Image.Image) -> str:
    variants = [
        preprocess_ocr_image(image),
        preprocess_ocr_image(image, threshold=155),
        preprocess_ocr_image(image, threshold=180),
    ]
    configs = [
        "--oem 3 --psm 6",
        "--oem 3 --psm 4",
        "--oem 3 --psm 11",
    ]
    candidates = []
    for variant in variants:
        for config in configs:
            text = pytesseract.image_to_string(variant, config=config)
            candidates.append((ocr_quality_score(text), clean_ocr_text(text)))
    best_score, best_text = max(candidates, key=lambda candidate: candidate[0])
    if best_score < 0.7:
        raise ValueError(
            "OCR text quality is too low. Upload a sharper, well-lit, straight image, "
            "or use a text-based PDF/DOCX for better questions."
        )
    return best_text

def split_into_chapters(text: str):
    # Match section headers like Chapter, Lesson, Unit with numbers or roman numerals
    pattern = re.compile(r"\b(Chapter|Lesson|Unit)\s+([0-9IVXLC]+)\b", re.IGNORECASE)
    matches = list(pattern.finditer(text))

    chapters = []
    for i, match in enumerate(matches):
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        section_type = match.group(1).capitalize()
        section_number = match.group(2)
        title = f"{section_type} {section_number}"
        content = text[start:end].strip()
        chapters.append({"title": title, "content": content, "section_type": section_type})

    return chapters

def split_into_study_sections(text: str, words_per_section: int = 900):
    """Create study sections when explicit chapter headings are not available."""
    words = text.split()
    if not words:
        return []

    sections = []
    for index, start in enumerate(range(0, len(words), words_per_section), start=1):
        section_words = words[start:start + words_per_section]
        content = " ".join(section_words).strip()
        if content:
            sections.append({
                "title": f"Study Section {index}",
                "content": content,
                "section_type": "Auto",
            })

    return sections

def build_chapters_from_text(text: str, document_hash: str = "edited"):
    cleaned = clean_ocr_text(text)
    chapters = split_into_chapters(cleaned)
    if not chapters or (len(chapters) == 1 and len(chapters[0].get("content", "")) > 6000):
        chapters = split_into_study_sections(cleaned, words_per_section=650)
    if not chapters:
        raise ValueError("No usable study text found. Add clearer text before applying changes.")
    for index, chapter in enumerate(chapters, start=1):
        chapter["hash"] = hash_content(f"{document_hash}:edited:{index}:{chapter['content']}")
    return chapters
